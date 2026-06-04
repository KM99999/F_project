"""Excel export (§5.7).

- Columns: fecha, cliente, monto, concepto, emisor, estado, score, link a imagen
  (+ código y nombre del carnet del cliente).
- Correct cell types: fecha as date, monto as number with currency-style format,
  image as a clickable hyperlink.
- A "Resumen" sheet with totals per estado.
"""

from datetime import date
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font

from app.db.models import Recibo

_ESTADO_LABEL = {
    "unico": "Único",
    "posible_duplicado": "Posible duplicado",
    "duplicado_confirmado": "Duplicado confirmado",
}

_HEADERS = [
    "ID", "Fecha", "Cliente", "Emisor", "Monto", "Moneda", "Concepto",
    "Estado", "Score", "Código carnet", "Nombre carnet", "Imagen",
]

_LINK_FONT = Font(color="0563C1", underline="single")
_BOLD = Font(bold=True)


def _as_date(value):
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except (ValueError, TypeError):
        return value


def build_workbook(recibos: Iterable[Recibo], media_base_url: str) -> bytes:
    base = (media_base_url or "").rstrip("/")
    wb = Workbook()

    ws = wb.active
    ws.title = "Recibos"
    ws.append(_HEADERS)
    for cell in ws[1]:
        cell.font = _BOLD

    rows = list(recibos)
    for r in rows:
        idx = ws.max_row + 1
        ws.cell(row=idx, column=1, value=r.id)

        fecha_val = _as_date(r.fecha)
        c_fecha = ws.cell(row=idx, column=2, value=fecha_val)
        if isinstance(fecha_val, date):
            c_fecha.number_format = "yyyy-mm-dd"

        ws.cell(row=idx, column=3, value=r.cliente_original)
        ws.cell(row=idx, column=4, value=r.emisor_original)

        c_monto = ws.cell(row=idx, column=5, value=float(r.monto) if r.monto is not None else None)
        c_monto.number_format = "#,##0.00"

        ws.cell(row=idx, column=6, value=r.moneda)
        ws.cell(row=idx, column=7, value=r.concepto)
        ws.cell(row=idx, column=8, value=_ESTADO_LABEL.get(r.estado, r.estado))
        ws.cell(row=idx, column=9, value=r.score)
        ws.cell(row=idx, column=10, value=r.carnet_codigo)
        ws.cell(row=idx, column=11, value=r.carnet_nombre)

        c_link = ws.cell(row=idx, column=12, value="Ver imagen" if r.imagen_url else None)
        if r.imagen_url:
            url = f"{base}{r.imagen_url}" if r.imagen_url.startswith("/") else r.imagen_url
            c_link.hyperlink = url
            c_link.font = _LINK_FONT

    widths = [6, 12, 24, 24, 12, 8, 28, 18, 7, 14, 24, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    # --- Resumen sheet: totals per estado ---
    resumen = wb.create_sheet("Resumen")
    resumen.append(["Estado", "Cantidad"])
    for cell in resumen[1]:
        cell.font = _BOLD
    counts = {"unico": 0, "posible_duplicado": 0, "duplicado_confirmado": 0}
    for r in rows:
        counts[r.estado] = counts.get(r.estado, 0) + 1
    for estado, label in _ESTADO_LABEL.items():
        resumen.append([label, counts.get(estado, 0)])
    resumen.append(["Total", len(rows)])
    resumen.cell(row=resumen.max_row, column=1).font = _BOLD
    resumen.cell(row=resumen.max_row, column=2).font = _BOLD
    resumen.column_dimensions["A"].width = 22
    resumen.column_dimensions["B"].width = 10

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
